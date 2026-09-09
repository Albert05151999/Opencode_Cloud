"""Exercise the actual formats used by the data Agent; fail on bad round trips."""
import json
from pathlib import Path
import numpy
import pandas
import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot
import openpyxl
import xlsxwriter
import pptx
import docx
import PIL.Image
import yaml
import markdown
import tabulate
import requests
import httpx

out = Path(__file__).resolve().parents[1] / 'artifacts/python-smoke'
out.mkdir(parents=True, exist_ok=True)
book = openpyxl.Workbook()
book.active.append(['value'])
book.active.append([42])
book.save(out / 'sample.xlsx')
assert pandas.read_excel(out / 'sample.xlsx')['value'].tolist() == [42]
slides = pptx.Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[0])
slide.shapes.title.text = 'Cloud Agent smoke'
slides.save(out / 'sample.pptx')
loaded = pptx.Presentation(out / 'sample.pptx')
assert len(loaded.slides) == 1 and loaded.slides[0].shapes.title.text == 'Cloud Agent smoke'
pyplot.plot(numpy.array([0, 1, 2]), [0, 1, 4])
pyplot.savefig(out / 'sample.png')
pyplot.close()
with PIL.Image.open(out / 'sample.png') as img:
    img.verify()
data = {'agent': 'agent-data', 'values': [1, 2, 3]}
(out / 'sample.yaml').write_text(yaml.safe_dump(data), encoding='utf-8')
assert yaml.safe_load((out / 'sample.yaml').read_text()) == data
report = {'result': 'passed', 'checks': ['imports', 'xlsx pandas roundtrip', 'pptx roundtrip', 'PNG decode', 'YAML roundtrip']}
(out / 'report.json').write_text(json.dumps(report, indent=2) + '\n')
print(json.dumps(report))
